#!/usr/bin/env python3
"""Turn a question bank or a sampled exam into something a student can use.

Markdown, CSV and a self-grading HTML page need nothing beyond the standard
library and are always available. Excel and Word need a third-party package
each, and follow the convention the anki deck exporter already set: the import
happens at the point of use and a missing package produces a sentence telling
you what to install, never a traceback.
"""

from __future__ import annotations

import csv
import html
import json
from dataclasses import dataclass
from pathlib import Path

from question_bank import CASE, MCQ, WRITTEN, BankQuestion, QuestionBank

OPTION_LETTERS = ("a", "b", "c", "d", "e")

COLUMNS = (
    "question_id",
    "kind",
    "lecture",
    "module",
    "stem",
    "option_a",
    "option_b",
    "option_c",
    "option_d",
    "correct_option",
    "answer",
    "years",
    "badges",
    "past_exam",
    "question_bank",
    "important",
    "repeats",
    "duplicate_of",
)


class ExportError(RuntimeError):
    """Raised when an export cannot be produced."""


@dataclass(frozen=True)
class ExportResult:
    path: Path
    written: int
    note: str = ""


def _row(bank: QuestionBank, question: BankQuestion) -> dict[str, object]:
    options = question.options
    return {
        "question_id": question.question_id,
        "kind": question.kind,
        "lecture": question.lecture,
        "module": question.module_id,
        "stem": question.stem,
        "option_a": options.get("a", ""),
        "option_b": options.get("b", ""),
        "option_c": options.get("c", ""),
        "option_d": options.get("d", ""),
        "correct_option": question.correct_option,
        "answer": question.answer,
        "years": ", ".join(str(year) for year in question.years),
        "badges": " | ".join(question.badge_labels),
        "past_exam": "yes" if question.is_past_exam else "",
        "question_bank": "yes" if question.is_question_bank else "",
        "important": "yes" if question.is_important else "",
        "repeats": bank.repeat_count(question),
        "duplicate_of": question.duplicate_of,
    }


def write_csv(
    bank: QuestionBank, path: Path | str, questions: tuple[BankQuestion, ...] | None = None
) -> ExportResult:
    """Write the bank as CSV. utf-8-sig so Excel opens Arabic correctly."""
    rows = questions if questions is not None else bank.questions
    target = Path(path)
    target.parent.mkdir(parents=True, exist_ok=True)
    with target.open("w", encoding="utf-8-sig", newline="") as handle:
        writer = csv.DictWriter(handle, fieldnames=list(COLUMNS))
        writer.writeheader()
        for question in rows:
            writer.writerow(_row(bank, question))
    return ExportResult(target, len(rows))


def write_json(
    bank: QuestionBank, path: Path | str, questions: tuple[BankQuestion, ...] | None = None
) -> ExportResult:
    rows = questions if questions is not None else bank.questions
    target = Path(path)
    target.parent.mkdir(parents=True, exist_ok=True)
    payload = {
        "module": bank.module_id,
        "lectures": list(bank.lectures),
        "unique": len(bank.unique),
        "total": len(bank.questions),
        "questions": [_row(bank, question) for question in rows],
    }
    target.write_text(
        json.dumps(payload, ensure_ascii=False, indent=2) + "\n", encoding="utf-8"
    )
    return ExportResult(target, len(rows))


def write_xlsx(
    bank: QuestionBank, path: Path | str, questions: tuple[BankQuestion, ...] | None = None
) -> ExportResult:
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Alignment, Font
        from openpyxl.utils import get_column_letter
    except ImportError as error:
        raise ExportError(
            "openpyxl is required for Excel export. Install it with "
            "`pip install openpyxl`, or export --format csv instead."
        ) from error

    rows = questions if questions is not None else bank.questions
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Questions"
    sheet.append(list(COLUMNS))
    for cell in sheet[1]:
        cell.font = Font(bold=True)
    sheet.freeze_panes = "A2"

    for question in rows:
        row = _row(bank, question)
        sheet.append([row[column] for column in COLUMNS])

    widths = {"stem": 60, "answer": 50, "lecture": 18, "question_id": 28}
    for index, column in enumerate(COLUMNS, start=1):
        sheet.column_dimensions[get_column_letter(index)].width = widths.get(column, 14)
    for row_cells in sheet.iter_rows(min_row=2):
        for cell in row_cells:
            cell.alignment = Alignment(vertical="top", wrap_text=True)
    sheet.auto_filter.ref = sheet.dimensions

    target = Path(path)
    target.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(target)
    return ExportResult(target, len(rows))


def write_docx(
    bank: QuestionBank,
    path: Path | str,
    questions: tuple[BankQuestion, ...],
    *,
    title: str = "Question paper",
    with_answers: bool = False,
) -> ExportResult:
    try:
        from docx import Document
    except ImportError as error:
        raise ExportError(
            "python-docx is required for Word export. Install it with "
            "`pip install python-docx`, or export --format md instead."
        ) from error

    document = Document()
    document.add_heading(title, level=1)
    for index, question in enumerate(questions, start=1):
        document.add_paragraph(f"{index}. {question.stem}")
        if question.scenario and question.kind == CASE:
            document.add_paragraph(question.scenario)
        for letter in OPTION_LETTERS:
            text = question.options.get(letter)
            if text:
                document.add_paragraph(f"{letter}. {text}", style="List Bullet")
        if with_answers:
            document.add_paragraph(f"Answer: {question.answer}")

    target = Path(path)
    target.parent.mkdir(parents=True, exist_ok=True)
    document.save(target)
    return ExportResult(target, len(questions))


def render_bank_markdown(bank: QuestionBank, questions: tuple[BankQuestion, ...] | None = None) -> str:
    rows = questions if questions is not None else bank.unique
    lines = [f"# {bank.module_id} — question bank", ""]
    lines.append(f"{len(bank.unique)} unique questions of {len(bank.questions)} collected.")
    lines.append("")
    for kind, heading in ((MCQ, "MCQs"), (WRITTEN, "Written Questions"), (CASE, "Clinical Cases")):
        of_kind = [question for question in rows if question.kind == kind]
        if not of_kind:
            continue
        lines.append(f"## {heading}")
        lines.append("")
        for index, question in enumerate(of_kind, start=1):
            badges = f" **[{', '.join(question.badge_labels)}]**" if question.badge_labels else ""
            repeats = bank.repeat_count(question)
            repeat_note = f" _(asked {repeats}×)_" if repeats > 1 else ""
            lines.append(f"### {index}. {question.stem}{badges}{repeat_note}")
            lines.append("")
            lines.append(f"_{question.lecture}_")
            lines.append("")
            for letter in OPTION_LETTERS:
                text = question.options.get(letter)
                if text:
                    lines.append(f"- **{letter}.** {text}")
            if question.options:
                lines.append("")
            if question.answer:
                lines.append(f"**Answer:** {question.answer}")
                lines.append("")
    return "\n".join(lines).rstrip() + "\n"


def render_exam_markdown(
    questions: tuple[BankQuestion, ...], *, title: str = "Exam paper"
) -> tuple[str, str]:
    """The paper and its answer key, as two documents.

    Separate on purpose: a paper with the answers printed under each question
    cannot be sat.
    """
    paper = [f"# {title}", "", f"{len(questions)} questions.", ""]
    key = [f"# {title} — answer key", ""]
    for index, question in enumerate(questions, start=1):
        if question.scenario:
            paper.append(f"**{index}.** {question.scenario}")
            paper.append("")
            for sub in question.sub_questions:
                paper.append(f"   {sub}")
        else:
            paper.append(f"**{index}.** {question.stem}")
        for letter in OPTION_LETTERS:
            text = question.options.get(letter)
            if text:
                paper.append(f"- {letter}. {text}")
        paper.append("")
        key.append(f"**{index}.** {question.answer or '—'}  _({question.lecture})_")
    return "\n".join(paper).rstrip() + "\n", "\n".join(key).rstrip() + "\n"


def render_exam_html(questions: tuple[BankQuestion, ...], *, title: str = "Exam paper") -> str:
    """A single self-contained page that marks itself.

    No dependency, no network, no build step: one file a student can open from
    a phone. Only MCQs can be auto-marked; written questions and cases show
    their model answer on reveal.
    """
    payload = [
        {
            "stem": question.stem,
            "scenario": question.scenario,
            "sub": list(question.sub_questions),
            "options": {letter: question.options[letter] for letter in OPTION_LETTERS if letter in question.options},
            "correct": question.correct_option,
            "answer": question.answer,
            "lecture": question.lecture,
            "years": list(question.years),
        }
        for question in questions
    ]
    data = json.dumps(payload, ensure_ascii=False)
    return f"""<!doctype html>
<html lang="en" dir="ltr">
<head>
<meta charset="utf-8">
<meta name="viewport" content="width=device-width, initial-scale=1">
<title>{html.escape(title)}</title>
<style>
  :root {{ color-scheme: light dark; --line: #d8d8d8; --ok: #1a7f37; --bad: #b42318; }}
  body {{ font: 16px/1.6 system-ui, -apple-system, "Segoe UI", sans-serif;
         max-width: 44rem; margin: 0 auto; padding: 1.5rem 1rem 6rem; }}
  h1 {{ font-size: 1.4rem; }}
  .q {{ border-top: 1px solid var(--line); padding: 1rem 0; }}
  .stem {{ font-weight: 600; }}
  .scenario {{ opacity: .85; margin: .4rem 0; }}
  label {{ display: block; padding: .4rem .6rem; border: 1px solid var(--line);
           border-radius: .4rem; margin: .3rem 0; cursor: pointer; }}
  label:hover {{ border-color: #999; }}
  .correct {{ border-color: var(--ok); }}
  .wrong {{ border-color: var(--bad); }}
  .model {{ display: none; background: rgba(127,127,127,.12); padding: .6rem;
            border-radius: .4rem; margin-top: .5rem; white-space: pre-wrap; }}
  .revealed .model {{ display: block; }}
  .meta {{ font-size: .82rem; opacity: .65; }}
  #bar {{ position: fixed; inset: auto 0 0 0; background: Canvas;
          border-top: 1px solid var(--line); padding: .7rem 1rem; text-align: center; }}
  button {{ font: inherit; padding: .45rem 1.1rem; border-radius: .4rem;
            border: 1px solid var(--line); background: Canvas; color: inherit; cursor: pointer; }}
</style>
</head>
<body>
<h1>{html.escape(title)}</h1>
<div id="paper"></div>
<div id="bar"><button id="mark">Mark my answers</button> <span id="score"></span></div>
<script>
const QUESTIONS = {data};
const paper = document.getElementById("paper");
QUESTIONS.forEach((q, i) => {{
  const box = document.createElement("section");
  box.className = "q";
  const stem = document.createElement("p");
  stem.className = "stem";
  stem.textContent = (i + 1) + ". " + q.stem;
  box.append(stem);
  if (q.scenario) {{
    const s = document.createElement("p");
    s.className = "scenario";
    s.textContent = q.scenario;
    box.append(s);
  }}
  (q.sub || []).forEach(line => {{
    const p = document.createElement("p");
    p.className = "scenario";
    p.textContent = line;
    box.append(p);
  }});
  const letters = Object.keys(q.options || {{}});
  letters.forEach(letter => {{
    const label = document.createElement("label");
    const input = document.createElement("input");
    input.type = "radio";
    input.name = "q" + i;
    input.value = letter;
    label.append(input, document.createTextNode(" " + letter + ". " + q.options[letter]));
    box.append(label);
  }});
  const model = document.createElement("div");
  model.className = "model";
  model.textContent = q.answer || "—";
  box.append(model);
  const meta = document.createElement("p");
  meta.className = "meta";
  meta.textContent = q.lecture + (q.years.length ? " · " + q.years.join(", ") : "");
  box.append(meta);
  paper.append(box);
}});
document.getElementById("mark").addEventListener("click", () => {{
  let right = 0, marked = 0;
  QUESTIONS.forEach((q, i) => {{
    const box = paper.children[i];
    box.classList.add("revealed");
    if (!q.correct) return;
    marked++;
    box.querySelectorAll("label").forEach(label => {{
      const input = label.querySelector("input");
      label.classList.remove("correct", "wrong");
      if (input.value === q.correct) label.classList.add("correct");
      else if (input.checked) label.classList.add("wrong");
    }});
    const chosen = box.querySelector("input:checked");
    if (chosen && chosen.value === q.correct) right++;
  }});
  document.getElementById("score").textContent =
    marked ? right + " / " + marked + " correct" : "model answers shown";
}});
</script>
</body>
</html>
"""


__all__ = [
    "COLUMNS",
    "ExportError",
    "ExportResult",
    "render_bank_markdown",
    "render_exam_html",
    "render_exam_markdown",
    "write_csv",
    "write_docx",
    "write_json",
    "write_xlsx",
]
